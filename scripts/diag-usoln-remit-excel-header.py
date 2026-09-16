"""Dump header + sample rows of usol_n weekly settlement excel.

Purpose: identify task_item_id mapping key + amount columns for snapshot backfill.
Run: python scripts/diag-usoln-remit-excel-header.py
"""
import glob
import os
import openpyxl

DOWNLOADS = os.path.expanduser("~/Downloads")
patterns = [
    "*7월 2주차 (1)*.xlsx",
    "*7월 2주차*.xlsx",
    "*주정산_7월*.xlsx",
]

candidates = []
for p in patterns:
    candidates.extend(glob.glob(os.path.join(DOWNLOADS, p)))

seen = set()
uniq = [c for c in candidates if not (c in seen or seen.add(c))]

if not uniq:
    print("No matching file found in", DOWNLOADS)
    for f in sorted(os.listdir(DOWNLOADS)):
        if "주정산" in f or "usol" in f.lower():
            print("  candidate:", f)
    raise SystemExit(1)

for path in uniq:
    print(f"\n=== {os.path.basename(path)} ===")
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"--- sheet: {sheet_name}  rows={ws.max_row} cols={ws.max_column} ---")
        # first 5 rows
        for r in range(1, min(6, ws.max_row + 1)):
            row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            print(f"  row{r}:", row)
        # last 2 rows (for total/summary detection)
        if ws.max_row > 6:
            print(f"  ...")
            for r in range(max(6, ws.max_row - 1), ws.max_row + 1):
                row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                print(f"  row{r}:", row)
